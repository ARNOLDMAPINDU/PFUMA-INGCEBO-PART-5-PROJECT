# -*- coding: utf-8 -*-
"""
thesis_builder.py — renders the TCL 5000 final project document to .docx in the
exact format the marking guidelines prescribe.

Guidelines implemented (see AGENT_BRIEF.md section 3):
  * Times New Roman throughout; body 12 pt; 1.5 line spacing
  * Left margin 1.5 in (binding), right/top/bottom 1 in
  * Chapter heading 14 pt BOLD UPPERCASE
  * Section heading 14 pt NORMAL UPPERCASE
  * Sub-section heading 13 pt bold; sub-sub-section 12 pt bold
  * Title page unnumbered; front matter lower-roman centre-bottom;
    body restarts at arabic 1 on Chapter 1
  * Table caption ABOVE the table, figure caption BELOW the figure, both 11 pt
  * Table entries 11 pt, single spaced
  * First paragraph after a heading flush left; subsequent paragraphs indent 0.5 in
  * References numbered by order of first appearance, resolved from [[key]] markers
  * TOC / List of Tables / List of Figures emitted as Word fields — open the file
    and press Ctrl+A then F9 to populate them

Usage:
    python thesis_builder.py                  # builds the full document
    python thesis_builder.py --out other.docx
"""
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))

FONT = "Times New Roman"
BODY_PT = 12
SMALL_PT = 11
INDENT = Inches(0.5)


# --------------------------------------------------------------------------
# low-level OXML helpers
# --------------------------------------------------------------------------
def _field(paragraph, instruction, placeholder=""):
    """Insert a Word field code (TOC, PAGE, ...) into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)
    return run


def _page_number_format(section, fmt, start=None):
    """Set the page-numbering format (lowerRoman / decimal) for a section."""
    sectPr = section._sectPr
    for existing in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(existing)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def _footer_page_number(section):
    """Centre-bottom page number in this section's footer."""
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.space_after = Pt(0)
    _field(para, "PAGE", "1")
    for run in para.runs:
        run.font.name = FONT
        run.font.size = Pt(BODY_PT)


def _blank_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for para in footer.paragraphs:
        para.text = ""


def _set_cell_font(cell, size_pt=SMALL_PT, bold=False):
    for para in cell.paragraphs:
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.first_line_indent = Pt(0)
        for run in para.runs:
            run.font.name = FONT
            run.font.size = Pt(size_pt)
            run.font.bold = bold


def _set_repeat_header(row):
    """Mark a table row as a header row that repeats across page breaks."""
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


# --------------------------------------------------------------------------
# styles
# --------------------------------------------------------------------------
def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    # Chapter = Heading 1 (14 pt bold UPPERCASE)
    # Section = Heading 2 (14 pt NORMAL UPPERCASE)  <- guidelines say normal, not bold
    # Sub-section = Heading 3 (13 pt bold)
    # Sub-sub-section = Heading 4 (12 pt bold)
    spec = {
        "Heading 1": (14, True),
        "Heading 2": (14, False),
        "Heading 3": (13, True),
        "Heading 4": (12, True),
    }
    for name, (size, bold) in spec.items():
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.keep_with_next = True

    # Separate caption styles so the List of Tables and List of Figures can be
    # generated as two independent Word fields keyed on style name.
    for cap_name in ("TableCaption", "FigureCaption", "ListingCaption"):
        if cap_name in [s.name for s in doc.styles]:
            continue
        cap = doc.styles.add_style(cap_name, 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH
        cap.base_style = doc.styles["Normal"]
        cap.font.name = FONT
        cap.font.size = Pt(SMALL_PT)
        cap.font.bold = False
        cap.font.italic = False
        cap.paragraph_format.line_spacing = 1.0
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.first_line_indent = Pt(0)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["TableCaption"].paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------
# citation resolution
# --------------------------------------------------------------------------
CITE_RE = re.compile(r"\[\[([A-Za-z0-9_.\-]+)\]\]")
# Defends against a caption already carrying its own "Figure 3.4:" or
# "Table 3.1:" prefix, which would otherwise double up with the prefix the
# renderer adds itself.
CAPTION_PREFIX_RE = re.compile(r"^(figure|table)\s+[\w.\-]+\s*[:.\-]\s*", re.I)


def _strip_caption_prefix(caption):
    return CAPTION_PREFIX_RE.sub("", caption, count=1)


class Citations:
    """Resolves [[key]] markers to [n] in order of first appearance."""

    def __init__(self, refkeys):
        self.refkeys = dict(refkeys)
        self.order = []
        self.unknown = set()

    def number(self, key):
        if key not in self.refkeys:
            self.unknown.add(key)
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def resolve(self, text):
        if not isinstance(text, str):
            return text

        def sub(match):
            return "[%d]" % self.number(match.group(1))

        return CITE_RE.sub(sub, text)

    def reference_list(self):
        return [(i + 1, self.refkeys.get(k, "MISSING REFERENCE FOR KEY: %s" % k))
                for i, k in enumerate(self.order)]


# --------------------------------------------------------------------------
# body rendering
# --------------------------------------------------------------------------
class Renderer:
    def __init__(self, doc, cites):
        self.doc = doc
        self.cites = cites
        self.tables = 0
        self.figures = 0
        self.missing_figures = []
        self._first_para_after_heading = True

    # -- headings ----------------------------------------------------------
    def heading(self, text, level):
        self.doc.add_heading(self.cites.resolve(text), level=level)
        self._first_para_after_heading = True

    # -- paragraphs --------------------------------------------------------
    def paragraph(self, text):
        para = self.doc.add_paragraph(self.cites.resolve(text))
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Guidelines: first paragraph after a title starts at the line start,
        # subsequent paragraphs are indented.
        para.paragraph_format.first_line_indent = (
            Pt(0) if self._first_para_after_heading else INDENT
        )
        self._first_para_after_heading = False
        return para

    def listing(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for item in items:
            para = self.doc.add_paragraph(self.cites.resolve(item), style=style)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        self._first_para_after_heading = False

    # -- tables ------------------------------------------------------------
    def table(self, spec):
        self.tables += 1
        number = spec.get("number") or str(self.tables)
        caption = _strip_caption_prefix(self.cites.resolve(spec.get("caption", "")))

        cap = self.doc.add_paragraph(style="TableCaption")
        run = cap.add_run("Table %s: %s" % (number, caption))
        run.font.name = FONT
        run.font.size = Pt(SMALL_PT)

        headers = spec["headers"]
        rows = spec.get("rows", [])
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for i, text in enumerate(headers):
            table.rows[0].cells[i].text = str(text)
            _set_cell_font(table.rows[0].cells[i], bold=True)
        _set_repeat_header(table.rows[0])

        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row[: len(headers)]):
                cells[i].text = self.cites.resolve(str(value))
                _set_cell_font(cells[i])

        self.doc.add_paragraph()
        self._first_para_after_heading = True

    # -- figures -----------------------------------------------------------
    def figure(self, spec):
        self.figures += 1
        number = spec.get("number") or str(self.figures)
        caption = _strip_caption_prefix(self.cites.resolve(spec.get("caption", "")))
        path = spec["path"]
        if not os.path.isabs(path):
            path = os.path.join(HERE, path)

        if os.path.exists(path):
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            para.add_run().add_picture(path, width=Inches(spec.get("width", 5.8)))
        else:
            self.missing_figures.append(path)
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("[ MISSING IMAGE: %s ]" % os.path.basename(path))
            run.bold = True

        cap = self.doc.add_paragraph(style="FigureCaption")
        run = cap.add_run("Figure %s: %s" % (number, caption))
        run.font.name = FONT
        run.font.size = Pt(SMALL_PT)
        self.doc.add_paragraph()
        self._first_para_after_heading = True

    # -- code listings -----------------------------------------------------
    def code(self, spec):
        """Render a source-code listing in a monospaced font.

        Either `text` is given directly, or `path` names a file to read
        (relative to the project folder). `max_lines` truncates long files with
        an explicit notice rather than silently dropping code.
        """
        caption = spec.get("caption")
        if caption:
            # ListingCaption, not TableCaption, so code listings do not appear
            # in the List of Tables.
            cap = self.doc.add_paragraph(style="ListingCaption")
            run = cap.add_run(caption)
            run.font.name = FONT
            run.font.size = Pt(SMALL_PT)
            run.bold = True

        truncated = 0
        total_lines = None
        if "text" in spec:
            body = spec["text"]
            lines = body.replace("\t", "    ").splitlines()
        else:
            path = spec["path"]
            if not os.path.isabs(path):
                path = os.path.join(HERE, path)
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as handle:
                    body = handle.read()
            except OSError as exc:
                body = "[ could not read %s: %s ]" % (path, exc)
            all_lines = body.replace("\t", "    ").splitlines()
            total_lines = len(all_lines)

            line_ranges = spec.get("line_ranges")
            shown_count = None
            if line_ranges:
                # Excerpt specific 1-indexed inclusive [start, end] ranges,
                # e.g. an enum/constants block plus one whole function,
                # rather than a naive head-of-file cut that could truncate
                # mid-function. A marker note is inserted between ranges.
                lines = []
                covered = 0
                for i, (start, end) in enumerate(line_ranges):
                    if i > 0:
                        lines.append(
                            "    /* ... %d lines omitted: unrelated setup, "
                            "command-parsing or telemetry code; see the "
                            "accompanying media ... */" % (start - prev_end - 1)
                        )
                    lines.extend(all_lines[start - 1:end])
                    covered += end - start + 1
                    prev_end = end
                shown_count = covered
                truncated = total_lines - covered
            else:
                lines = all_lines
                max_lines = spec.get("max_lines")
                if max_lines and len(lines) > max_lines:
                    truncated = len(lines) - max_lines
                    lines = lines[:max_lines]
                shown_count = len(lines)

        for line in lines:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.line_spacing = 1.0
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.first_line_indent = Pt(0)
            pf.left_indent = Inches(0.25)
            run = para.add_run(line if line.strip() else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")

        if truncated:
            para = self.doc.add_paragraph()
            para.paragraph_format.first_line_indent = Pt(0)
            note = (
                "[ %d of %d lines shown above; %d lines omitted from this "
                "excerpt. The complete source is supplied on the "
                "accompanying media. ]" % (shown_count, total_lines, truncated)
                if total_lines is not None
                else "[ listing truncated here: %d further lines omitted. "
                     "The complete source is supplied on the accompanying "
                     "media. ]" % truncated
            )
            run = para.add_run(note)
            run.italic = True
            run.font.size = Pt(SMALL_PT)

        self.doc.add_paragraph()
        self._first_para_after_heading = True

    # -- dispatch ----------------------------------------------------------
    def item(self, item):
        if isinstance(item, str):
            self.paragraph(item)
        elif isinstance(item, dict):
            if "subheading" in item:
                self.heading(item["subheading"], 3)
            elif "subsubheading" in item:
                self.heading(item["subsubheading"], 4)
            elif "bullets" in item:
                self.listing(item["bullets"], numbered=False)
            elif "numbered" in item:
                self.listing(item["numbered"], numbered=True)
            elif "table" in item:
                self.table(item["table"])
            elif "figure" in item:
                self.figure(item["figure"])
            elif "code" in item:
                self.code(item["code"])
            elif "pagebreak" in item:
                self.doc.add_page_break()
            else:
                raise ValueError("Unrecognised body item keys: %r" % sorted(item))
        else:
            raise TypeError("Unrecognised body item type: %r" % type(item))


# --------------------------------------------------------------------------
# front matter
# --------------------------------------------------------------------------
def centred(doc, text, size=BODY_PT, bold=False, italic=False, spacing=1.5):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.line_spacing = spacing
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return para


def _spacer(doc, count=1, pt=6):
    for _ in range(count):
        para = doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(pt)


def add_title_page(doc, tp):
    """Guidelines sizes: title 18-20 pt bold UPPER CASE, name/supervisor/
    statement/date 14 pt bold. Spacing is kept tight so the whole page fits on
    one sheet - an overflow would create an unnumbered orphan page."""
    centred(doc, tp["programme_banner"], size=14, bold=True, spacing=1.15)
    _spacer(doc)
    centred(doc, tp["institution"], size=13, bold=True, spacing=1.0)
    centred(doc, tp["faculty"], size=12, spacing=1.0)
    centred(doc, tp["department"], size=12, spacing=1.0)
    _spacer(doc)
    centred(doc, "PROJECT TITLE:", size=14, bold=True, spacing=1.0)
    centred(doc, tp["title"].upper(), size=18, bold=True, spacing=1.15)
    _spacer(doc)
    centred(doc, "By", size=14, bold=True, spacing=1.0)
    _spacer(doc)
    centred(doc, "Student Name: %s" % tp["student"], size=14, bold=True, spacing=1.0)
    centred(doc, "Student Number: %s" % tp["reg_number"], size=14, bold=True, spacing=1.0)
    centred(doc, "Supervised by: %s" % tp["supervisor"], size=14, bold=True, spacing=1.0)
    _spacer(doc)
    centred(doc, tp["submission_statement"], size=14, bold=True, spacing=1.15)
    _spacer(doc)
    centred(doc, "Submission Date: %s" % tp["date"], size=14, bold=True, spacing=1.0)


def add_committee_report(doc, text):
    doc.add_heading("EXAMINING COMMITTEE REPORT", level=1)
    para = doc.add_paragraph(text)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    labels = [
        ["Supervisor:", "Examiner 1:", "Examiner 2:"],
        ["Name: ..........................", "Name: ..........................",
         "Name: .........................."],
        ["Signature: ....................", "Signature: ....................",
         "Signature: ...................."],
        ["Date:     /     /", "Date:     /     /", "Date:     /     /"],
    ]
    for r, row in enumerate(labels):
        for c, text_ in enumerate(row):
            table.rows[r].cells[c].text = text_
            _set_cell_font(table.rows[r].cells[c], bold=(r == 0))
    doc.add_page_break()


def add_simple_page(doc, title, content):
    doc.add_heading(title, level=1)
    paragraphs = [content] if isinstance(content, str) else list(content)
    for text in paragraphs:
        para = doc.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Pt(0)
    doc.add_page_break()


def add_abstract(doc, fm):
    doc.add_heading("ABSTRACT", level=1)
    for text in fm.ABSTRACT:
        para = doc.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Pt(0)
    keywords = getattr(fm, "KEYWORDS", None)
    if keywords:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(0)
        run = para.add_run("Keywords: ")
        run.bold = True
        para.add_run(", ".join(keywords))
    doc.add_page_break()


def add_field_page(doc, title, instruction):
    doc.add_heading(title, level=1)
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    _field(para, instruction,
           "Press Ctrl+A then F9 in Word to populate this list.")
    doc.add_page_break()


def add_acronyms(doc, acronyms):
    doc.add_heading("LIST OF ACRONYMS AND SYMBOLS", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Acronym"
    table.rows[0].cells[1].text = "Expansion"
    for cell in table.rows[0].cells:
        _set_cell_font(cell, bold=True)
    _set_repeat_header(table.rows[0])
    for acronym, expansion in acronyms:
        cells = table.add_row().cells
        cells[0].text = acronym
        cells[1].text = expansion
        for cell in cells:
            _set_cell_font(cell)
    doc.add_page_break()


# --------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------
def load_parts():
    sys.path.insert(0, HERE)
    from chapters import frontmatter as fm  # noqa: E402

    chapters = []
    refkeys = {}
    for name in ("ch1", "ch2", "ch3", "ch4", "ch5"):
        module = __import__("chapters.%s" % name, fromlist=["chapters"])
        chapters.append(module)
        refkeys.update(getattr(module, "REFKEYS", {}))
    refkeys.update(getattr(fm, "REFKEYS", {}))
    return fm, chapters, refkeys


def build(out_path):
    fm, chapters, refkeys = load_parts()
    cites = Citations(refkeys)

    doc = Document()
    configure_styles(doc)

    # ---- section 1: title page, unnumbered --------------------------------
    sec1 = doc.sections[0]
    for sec in (sec1,):
        sec.left_margin = Inches(1.5)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
    _blank_footer(sec1)
    add_title_page(doc, fm.TITLE_PAGE)

    # ---- section 2: front matter, lower roman -----------------------------
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.left_margin, sec2.right_margin = Inches(1.5), Inches(1.0)
    sec2.top_margin, sec2.bottom_margin = Inches(1.0), Inches(1.0)
    _page_number_format(sec2, "lowerRoman", start=1)
    _footer_page_number(sec2)

    add_committee_report(doc, fm.COMMITTEE_REPORT)
    if getattr(fm, "DEDICATION", None):
        add_simple_page(doc, "DEDICATION", fm.DEDICATION)
    if getattr(fm, "ACKNOWLEDGEMENTS", None):
        add_simple_page(doc, "ACKNOWLEDGEMENTS", fm.ACKNOWLEDGEMENTS)
    add_simple_page(doc, "COPYRIGHT STATEMENT", fm.COPYRIGHT)
    add_simple_page(doc, "DECLARATION", fm.DECLARATION)
    add_abstract(doc, fm)
    add_field_page(doc, "TABLE OF CONTENTS", 'TOC \\o "1-3" \\h \\z \\u')
    add_field_page(doc, "LIST OF TABLES", 'TOC \\h \\z \\t "TableCaption,1"')
    add_field_page(doc, "LIST OF FIGURES", 'TOC \\h \\z \\t "FigureCaption,1"')
    add_acronyms(doc, fm.ACRONYMS)

    # ---- section 3: body, arabic restarting at 1 --------------------------
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec3.left_margin, sec3.right_margin = Inches(1.5), Inches(1.0)
    sec3.top_margin, sec3.bottom_margin = Inches(1.0), Inches(1.0)
    _page_number_format(sec3, "decimal", start=1)
    _footer_page_number(sec3)

    renderer = Renderer(doc, cites)
    for index, chapter in enumerate(chapters):
        renderer.heading(chapter.HEADING, 1)
        for section in chapter.SECTIONS:
            renderer.heading(section["heading"], 2)
            for item in section["body"]:
                renderer.item(item)
        if index != len(chapters) - 1:
            doc.add_page_break()

    # ---- references -------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("REFERENCES", level=1)
    references = cites.reference_list()
    for number, text in references:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.left_indent = Inches(0.5)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run("[%d]  " % number)
        run.bold = True
        para.add_run(text)

    # ---- appendices -------------------------------------------------------
    appendices = []
    try:
        from chapters import appendices as appendix_module
        appendices = getattr(appendix_module, "APPENDICES", [])
    except ImportError:
        pass

    if appendices:
        doc.add_page_break()
        doc.add_heading("APPENDICES", level=1)
        para = doc.add_paragraph(
            "The appendices are supplied in fulfilment of the requirement for "
            "system documentation, user documentation, simulation results and a "
            "code listing. They do not count towards the body page limit."
        )
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for appendix in appendices:
            doc.add_page_break()
            renderer.heading(
                "%s: %s" % (appendix["label"], appendix["title"]), 1
            )
            for item in appendix["body"]:
                renderer.item(item)

    doc.save(out_path)

    # ---- build report -----------------------------------------------------
    print("Saved: %s" % out_path)
    print("  chapters   : %d" % len(chapters))
    print("  tables     : %d" % renderer.tables)
    print("  figures    : %d" % renderer.figures)
    print("  references : %d" % len(references))
    if cites.unknown:
        print("  !! UNRESOLVED CITATION KEYS (no REFKEYS entry):")
        for key in sorted(cites.unknown):
            print("       [[%s]]" % key)
    if renderer.missing_figures:
        print("  !! MISSING IMAGE FILES:")
        for path in renderer.missing_figures:
            print("       %s" % path)
    if not cites.unknown and not renderer.missing_figures:
        print("  no unresolved citations, no missing images")
    print("\nOpen in Word and press Ctrl+A then F9 to populate the")
    print("Table of Contents, List of Tables and List of Figures.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--out",
        default=os.path.join(
            HERE,
            "PFUMA_INGCEBO - Project Documentation.docx",
        ),
    )
    args = parser.parse_args()
    build(args.out)
